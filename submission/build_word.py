from pathlib import Path
import argparse, copy, re, shutil, subprocess, tempfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission'
parser=argparse.ArgumentParser(description='Build the editable Word copy from the LaTeX manuscript.')
parser.add_argument('--pandoc', help='Path to Pandoc 3.x; defaults to pandoc on PATH')
args=parser.parse_args()
s=(OUT/'relay_uas_aeroconf.tex').read_text(encoding='utf-8-sig')
paper_title=re.search(r'\\title\{([^}]+)\}',s).group(1)
keys=re.findall(r'\\bibitem\{([^}]+)\}',s)
for i,k in enumerate(keys,1): s=s.replace('\\cite{'+k+'}',f'[{i}]')
labels={}
for kind in ('figure','table'):
    blocks=re.findall(r'\\begin\{'+kind+r'\*?\}.*?\\end\{'+kind+r'\*?\}',s,re.S)
    for number,block in enumerate(blocks,1):
        match=re.search(r'\\label\{([^}]+)\}',block)
        if match: labels[match.group(1)]=str(number)
for number,match in enumerate(re.finditer(r'\\section\{[^}]+\}\\label\{([^}]+)\}',s),1):
    labels[match.group(1)]=str(number)
for key,number in labels.items(): s=s.replace('\\ref{'+key+'}',number)
unresolved=re.findall(r'\\ref\{([^}]+)\}',s)
if unresolved: raise ValueError(f'Unresolved LaTeX references: {unresolved}')
s=s[s.index('\\begin{abstract}')+len('\\begin{abstract}'):]
s=s.replace('\\end{abstract}', '\n\\section*{CONTENTS_FIELD}\n')
s=s.replace('\\tableofcontents','')
s=s.replace('\\acknowledgements',r'\section*{Acknowledgements}')
s=s.replace('\\thebiography',r'\section*{Biography}')
bio=re.search(r'\\begin\{biographywithpic\}\s*\{([^}]+)\}\{([^}]+)\}\s*(.*?)\\end\{biographywithpic\}',s,re.S)
if bio is None: raise ValueError('Expected biographywithpic block not found')
biography_photo=OUT/bio.group(2)
if not biography_photo.is_file(): raise FileNotFoundError(biography_photo)
s=s.replace(bio.group(0),'\nBIOGRAPHY_PHOTO_MARKER\n\n\\textbf{'+bio.group(1)+'} '+bio.group(3).strip()+'\n')
s=re.sub(r'^%.*$', '', s, flags=re.M)
s=re.sub(r'\\begin\{minipage\}\{[^}]+\}', '', s)
s=s.replace('\\end{minipage}','').replace('\\footnotesize','').replace('\\small','')
s=re.sub(r'\\vspace\{[^}]+\}', '', s)
s=re.sub(r'\\gate\{([^}]+)\}',r'\1',s)
s=s.replace('\\appendices','').replace('\\FloatBarrier','')
s=re.sub(r'\{\\rm ([^}]+)\}',r'{\\mathrm{\1}}',s)
secnum=0
def section(m):
    global secnum
    title=m.group(1)
    if title=='Complete Carrier Mass Balance':return '\\section*{Appendix A. Complete Carrier Mass Balance}'
    secnum+=1
    return '\\section*{'+str(secnum)+'. '+title+'}'
s=re.sub(r'\\section\{([^}]+)\}',section,s)
figures=[]
def figure(m):
    x=m.group(0);name=re.search(r'figs/([^}]+)\.pdf',x).group(1)
    cap=re.search(r'\\caption\{(.*?)\}\s*\\label',x,re.S).group(1)
    key=re.search(r'\\label\{([^}]+)\}',x).group(1);num=labels.get(key,str(len(figures)+1))
    figures.append((name,cap,num,'figure*' in x))
    return '\nFIGUREMARKER'+str(len(figures)-1)+'\n\n'
s=re.sub(r'\\begin\{figure\*?\}.*?\\end\{figure\*?\}',figure,s,flags=re.S)
def table(m):
    x=m.group(0);cap=re.search(r'\\caption\{(.*?)\}',x,re.S).group(1)
    key=re.search(r'\\label\{([^}]+)\}',x).group(1)
    x=re.sub(r'\\caption\{.*?\}', '', x, count=1,flags=re.S)
    x=re.sub(r'\\begin\{table\*?\}(\[.*?\])?', '', x)
    x=re.sub(r'\\end\{table\*?\}', '', x)
    x=re.sub(r'\\setlength\{[^}]+\}\{[^}]+\}', '',x)
    x=re.sub(r'\\begin\{tabular\}\{.*?\}\s*\\toprule',lambda _: '\\begin{tabular}{'+'l'*(x[x.index('\\toprule'):].split('\\\\')[0].count('&')+1)+'}\n\\toprule',x, count=1,flags=re.S)
    return '\n\\textbf{Table '+labels.get(key,'?')+'. '+cap+'}\n\n'+x
s=re.sub(r'\\begin\{table\*?\}.*?\\end\{table\*?\}',table,s,flags=re.S)
s=re.sub(r'\\label\{[^}]+\}','',s)
s=re.sub(r'\\begin\{thebibliography\}\{[^}]+\}', r'\\section*{References}',s)
s=re.sub(r'\\setlength\{\\itemsep\}\{[^}]+\}','',s)
for i,k in enumerate(keys,1):s=s.replace('\\bibitem{'+k+'}',f'\n[{i}] ')
s=s.replace('\\end{thebibliography}','').replace('\\end{document}','')
s=s.replace('\\nolinkurl','\\texttt')
eqnum=0
def eq(m):
    global eqnum
    eqnum+=1
    return m.group(0).replace('\\end{equation}',r'\qquad\text{('+str(eqnum)+r')}\end{equation}')
s=re.sub(r'\\begin\{equation\}.*?\\end\{equation\}',eq,s,flags=re.S)
work=tempfile.TemporaryDirectory(prefix='relay-word-')
conv=Path(work.name)/'word-source.tex';conv.write_text(s,encoding='utf-8')
raw=Path(work.name)/'raw.docx'
pandoc=args.pandoc or shutil.which('pandoc')
if pandoc is None: raise SystemExit('Pandoc 3.x is required; install it on PATH or pass --pandoc PATH')
version=subprocess.run([str(pandoc),'--version'],capture_output=True,text=True,check=True).stdout
major=int(re.search(r'pandoc (\d+)\.',version).group(1))
if major<3: raise SystemExit('Pandoc 3.x is required')
subprocess.run([str(pandoc),str(conv),'-f','latex','-t','docx','-o',str(raw)],check=True)
doc=Document(raw)
sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
for k in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,k,Inches(.75))
sec.footer_distance=Inches(.3)
for st in doc.styles:
    if st.type==1:
        st.font.name='Times New Roman';st.font.size=Pt(10);st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.space_after=Pt(10)
        st.paragraph_format.line_spacing=1
        st.paragraph_format.first_line_indent=Pt(0)
for name in ['Heading 1','Heading 2']:
    st=doc.styles[name];st.font.size=Pt(12 if name=='Heading 1' else 10);st.font.bold=(name=='Heading 1');st.font.italic=(name=='Heading 2')
    st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER if name=='Heading 1' else WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5);st.paragraph_format.keep_with_next=True
for p in doc.paragraphs:
    p.paragraph_format.widow_control=True
    if p.style.name not in ['Heading 1','Heading 2']:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
# Native continuous section endings specify the column layout of preceding content.
def section_end(before,cols):
    p=OxmlElement('w:p');pp=OxmlElement('w:pPr');sp=copy.deepcopy(sec._sectPr)
    typ=sp.find(qn('w:type'))
    if typ is None:typ=OxmlElement('w:type');sp.insert(0,typ)
    typ.set(qn('w:val'),'continuous')
    co=sp.find(qn('w:cols'))
    if co is None:co=OxmlElement('w:cols');sp.append(co)
    co.set(qn('w:num'),str(cols));co.set(qn('w:space'),'360')
    pp.append(sp);p.append(pp);before.addprevious(p)
    return p
first=doc.paragraphs[0]
title=first.insert_paragraph_before(paper_title,style='Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after=Pt(10)
for r in title.runs:r.font.size=Pt(20);r.font.name='Times New Roman';r.bold=True
au=first.insert_paragraph_before('Noah Stephenson\nUnited States Military Academy\nWest Point, NY 10996\nnoah.stephenson@westpoint.edu')
au.alignment=WD_ALIGN_PARAGRAPH.CENTER
for r in au.runs:r.bold=True
section_end(first._p,1)
first.text='Abstract: '+first.text
first.paragraph_format.space_after=Pt(10)
for r in first.runs:r.bold=True;r.font.size=Pt(9)
for p in list(doc.paragraphs):
    if p.text=='CONTENTS_FIELD':
        p.text='Table of Contents';p.style='TOC Heading'
        f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'TOC \\o "1-1" \\t "Unnumbered Section,1" \\h \\z')
        np=OxmlElement('w:p');np.append(f);p._p.addnext(np)
    if p.text=='BIOGRAPHY_PHOTO_MARKER':
        p.text=''
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.add_run().add_picture(str(biography_photo),width=Inches(1.25))
        p.paragraph_format.keep_with_next=True
    if p.text.startswith('FIGUREMARKER'):
        idx=int(p.text.replace('FIGUREMARKER',''));name,cap,num,wide=figures[idx]
        p.text=''
        word_column_figure = name == 'fig3_service_map'
        if wide and not word_column_figure:section_end(p._p,2)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        image_name = name + ('_word' if word_column_figure else '') + '.png'
        picture_width = 7 if wide and not word_column_figure else 3.375
        p.add_run().add_picture(str(OUT/'figs'/image_name),width=Inches(picture_width))
        p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
        if word_column_figure:
            cap=cap.replace('pass means','P means').replace('size marks','S marks').replace('link fail marks','L marks')
        cp=doc.add_paragraph('Figure '+num+'. '+cap);p._p.addnext(cp._p)
        cp.alignment=WD_ALIGN_PARAGRAPH.CENTER;cp.paragraph_format.space_after=Pt(10)
        for r in cp.runs:r.bold=True;r.font.size=Pt(10)
        if wide and not word_column_figure:
            nx=cp._p.getnext()
            if nx is not None:section_end(nx,1)
    if p.text.startswith('Table '):
        p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:r.bold=True
for t in doc.tables:
    wide=len(t.columns)>6
    width=7 if wide else 3.375
    if wide:
        caption=t._tbl.getprevious();section_end(caption,2)
        nxt=t._tbl.getnext()
        if nxt is not None:section_end(nxt,1)
    t.autofit=False
    widths=[.50,.50,.50,.60,.625,.65] if len(t.columns)==6 else [width/len(t.columns)]*len(t.columns)
    for col,cw in zip(t.columns,widths):col.width=Inches(cw)
    for row in t.rows:
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for cell,cw in zip(row.cells,widths):
            cell.width=Inches(cw)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
                for r in p.runs:r.font.size=Pt(9)
    hdr=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(hdr)
    for cell in t.rows[0].cells:
        for p in cell.paragraphs:
            for r in p.runs:r.bold=True
co=sec._sectPr.find(qn('w:cols'))
if co is None:co=OxmlElement('w:cols');sec._sectPr.append(co)
co.set(qn('w:num'),'2');co.set(qn('w:space'),'360')
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
# Copyright is an unnumbered first-page footnote, retained from the source.
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT
xml=b'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote><w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:pPr><w:spacing w:after="0"/></w:pPr><w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t>U.S. Government work not protected by U.S. copyright</w:t></w:r></w:p></w:footnote></w:footnotes>'''
part=Part(PackURI('/word/footnotes.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',xml,doc.part.package)
existing=[r.target_part for r in doc.part.rels.values() if r.reltype==RT.FOOTNOTES]
if existing:existing[0]._blob=xml
else:doc.part.relate_to(part,RT.FOOTNOTES)
r=OxmlElement('w:r');fn=OxmlElement('w:footnoteReference');fn.set(qn('w:id'),'1');fn.set(qn('w:customMarkFollows'),'1');r.append(fn);first._p.append(r)
doc.core_properties.author='Noah Stephenson';doc.core_properties.title=title.text
# Remove theme-font overrides that can supersede Times New Roman.
for root in [doc.styles.element,doc.element]:
 for rf in root.xpath('//w:rFonts'):
  for a in list(rf.attrib):
   if 'Theme' in a:del rf.attrib[a]
  rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
un=doc.styles.add_style('Unnumbered Section',1);un.base_style=doc.styles['Normal'];un.font.name='Times New Roman';un.font.size=Pt(12);un.font.bold=True;un.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER;un.paragraph_format.space_before=Pt(10);un.paragraph_format.keep_with_next=True
for p in doc.paragraphs:
 if p.text in ['Data Availability','Acknowledgements','References','Biography']:
  p.style=un
  pp=p._p.get_or_add_pPr();ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');pp.append(ol)
out=OUT/'relay_uas_aeroconf.docx';doc.save(out)
work.cleanup()
print('Wrote',out,'equations',eqnum,'tables',len(doc.tables),'figures',len(figures),'plus biography photo')
